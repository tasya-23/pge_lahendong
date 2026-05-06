import io
import urllib.parse
from flask import Blueprint, jsonify, request, make_response
from helpers import (load_absensi, filter_df,
                     hitung_status, hitung_tidak_hadir, hitung_jam_kerja)
from middleware import login_required

export_bp = Blueprint("export", __name__)


# ==============================
# 📥 EXPORT EXCEL
# ==============================
@export_bp.route("/export/excel")
@login_required
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    df = load_absensi()
    if df is None:
        return "Tidak ada data", 404

    df = filter_df(df, request.args)

    # Bangun rekap per karyawan
    rows = []
    for pid, grp in df.groupby("Person ID"):
        name   = grp["Name"].iloc[0] if "Name" in grp.columns else str(pid)
        divisi = grp["Divisi"].iloc[0] if "Divisi" in grp.columns else "-"
        hadir, terlambat = hitung_status(grp)
        tidak_hadir      = hitung_tidak_hadir(grp)
        jam_kerja        = hitung_jam_kerja(grp)
        avg_jam          = round(jam_kerja / hadir, 1) if hadir > 0 else 0
        rows.append({
            "ID": str(pid), "Nama": name, "Divisi": divisi,
            "Hadir": hadir, "Terlambat": terlambat,
            "Tidak Hadir": tidak_hadir,
            "Total Jam": jam_kerja, "Avg Jam/Hari": avg_jam,
        })
    rows.sort(key=lambda x: x["Hadir"], reverse=True)

    # Warna
    GREEN  = "1A5C3A"
    WHITE  = "FFFFFF"
    GLIGHT = "E8F5EE"
    YELLOW = "FFF8E6"
    RED    = "FFF1F2"

    def fill(hex_color):
        return PatternFill("solid", fgColor=hex_color)

    def bdr():
        s = Side(style="thin", color="E5E7EB")
        return Border(left=s, right=s, top=s, bottom=s)

    # Buat workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Rekap Absensi"

    # Judul
    ws.merge_cells("A1:I1")
    ws["A1"] = "REKAP ABSENSI KARYAWAN — PGE LAHENDONG"
    ws["A1"].font      = Font(bold=True, size=13, color=WHITE)
    ws["A1"].fill      = fill(GREEN)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Sub-judul periode
    BULAN_NAMA = ["","Januari","Februari","Maret","April","Mei","Juni",
                  "Juli","Agustus","September","Oktober","November","Desember"]
    b_idx = request.args.get("bulan", "0")
    tahun = request.args.get("tahun", "0")
    divisi_filter = request.args.get("divisi", "") or "Semua Divisi"
    try:
        b_nama = BULAN_NAMA[int(b_idx)] if b_idx != "0" else "Semua Bulan"
    except Exception:
        b_nama = "Semua Bulan"
    t_nama = tahun if tahun != "0" else ""

    ws.merge_cells("A2:I2")
    ws["A2"] = f"Periode: {b_nama} {t_nama}  |  Divisi: {divisi_filter}"
    ws["A2"].font      = Font(size=10, color="374151")
    ws["A2"].fill      = fill("F3F4F6")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # Header kolom
    headers    = ["No","ID","Nama","Divisi","Hadir","Terlambat","Tidak Hadir","Total Jam","Avg Jam/Hari"]
    col_widths = [5, 14, 28, 18, 9, 11, 13, 12, 13]
    for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
        c = ws.cell(row=3, column=ci, value=h)
        c.font      = Font(bold=True, size=10, color=WHITE)
        c.fill      = fill(GREEN)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = bdr()
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[3].height = 20

    # Baris data
    for ri, row in enumerate(rows, 1):
        r_xl = ri + 3
        bg   = GLIGHT if ri % 2 == 0 else WHITE
        vals = [ri, row["ID"], row["Nama"], row["Divisi"],
                row["Hadir"], row["Terlambat"], row["Tidak Hadir"],
                row["Total Jam"], row["Avg Jam/Hari"]]
        for ci, val in enumerate(vals, 1):
            c = ws.cell(row=r_xl, column=ci, value=val)
            c.border    = bdr()
            c.alignment = Alignment(
                horizontal="left" if ci == 3 else "center",
                vertical="center"
            )
            if ci == 5 and isinstance(val, int):
                c.fill = fill(GLIGHT)
                c.font = Font(size=10, bold=True, color=GREEN)
            elif ci == 6 and isinstance(val, int) and val > 0:
                c.fill = fill(YELLOW)
                c.font = Font(size=10, bold=True, color="A16207")
            elif ci == 7 and isinstance(val, int) and val > 0:
                c.fill = fill(RED)
                c.font = Font(size=10, bold=True, color="DC3545")
            else:
                c.fill = fill(bg)
                c.font = Font(size=10)
        ws.row_dimensions[r_xl].height = 17

    # Baris total
    tr = len(rows) + 4
    ws.merge_cells(f"A{tr}:D{tr}")
    for ci in range(1, 10):
        c = ws.cell(row=tr, column=ci)
        c.fill      = fill(GREEN)
        c.font      = Font(bold=True, size=10, color=WHITE)
        c.border    = bdr()
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=tr, column=1, value="TOTAL")
    totals = {
        5: sum(r["Hadir"]        for r in rows),
        6: sum(r["Terlambat"]    for r in rows),
        7: sum(r["Tidak Hadir"]  for r in rows),
        8: round(sum(r["Total Jam"] for r in rows), 1),
    }
    totals[9] = round(totals[8] / max(totals[5], 1), 1)
    for ci, val in totals.items():
        ws.cell(row=tr, column=ci, value=val)
    ws.row_dimensions[tr].height = 20
    ws.freeze_panes = "A4"

    # Kirim file
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = urllib.parse.quote(f"Rekap_Absensi_{b_nama}_{t_nama}.xlsx")
    resp  = make_response(buf.read())
    resp.headers["Content-Type"] = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    resp.headers["Content-Disposition"] = f"attachment; filename*=UTF-8''{fname}"
    return resp


# ==============================
# 📄 EXPORT PDF
# ==============================
@export_bp.route("/export/pdf")
@login_required
def export_pdf():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import (SimpleDocTemplate, Table,
                                    TableStyle, Paragraph, Spacer)
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    df = load_absensi()
    if df is None:
        return "Tidak ada data", 404

    df = filter_df(df, request.args)

    rows = []
    for pid, grp in df.groupby("Person ID"):
        name   = grp["Name"].iloc[0] if "Name" in grp.columns else str(pid)
        divisi = grp["Divisi"].iloc[0] if "Divisi" in grp.columns else "-"
        hadir, terlambat = hitung_status(grp)
        tidak_hadir      = hitung_tidak_hadir(grp)
        jam_kerja        = hitung_jam_kerja(grp)
        rows.append([name, divisi, hadir, terlambat, tidak_hadir, jam_kerja])
    rows.sort(key=lambda x: x[2], reverse=True)

    BULAN_NAMA = ["","Januari","Februari","Maret","April","Mei","Juni",
                  "Juli","Agustus","September","Oktober","November","Desember"]
    b_idx = request.args.get("bulan", "0")
    tahun = request.args.get("tahun", "0")
    divisi_filter = request.args.get("divisi", "") or "Semua Divisi"
    try:
        b_nama = BULAN_NAMA[int(b_idx)] if b_idx != "0" else "Semua Bulan"
    except Exception:
        b_nama = "Semua Bulan"
    t_nama = tahun if tahun != "0" else ""

    GREEN   = colors.HexColor("#1A5C3A")
    GREEN_L = colors.HexColor("#E8F5EE")
    YELLOW  = colors.HexColor("#FFF8E6")
    RED_L   = colors.HexColor("#FFF1F2")
    GREY_L  = colors.HexColor("#F3F4F6")
    WHITE   = colors.white

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=15*mm, rightMargin=15*mm,
                            topMargin=15*mm, bottomMargin=15*mm)

    title_style = ParagraphStyle("t", fontSize=13, fontName="Helvetica-Bold",
                                 textColor=WHITE, alignment=TA_CENTER)
    sub_style   = ParagraphStyle("s", fontSize=9, fontName="Helvetica",
                                 textColor=colors.HexColor("#374151"),
                                 alignment=TA_CENTER)
    elems = []

    # Judul
    title_tbl = Table(
        [[Paragraph("REKAP ABSENSI KARYAWAN — PGE LAHENDONG", title_style)]],
        colWidths=[267*mm]
    )
    title_tbl.setStyle(TableStyle([
        ("BACKGROUND",    (0,0), (-1,-1), GREEN),
        ("TOPPADDING",    (0,0), (-1,-1), 8),
        ("BOTTOMPADDING", (0,0), (-1,-1), 8),
    ]))
    elems.append(title_tbl)
    elems.append(Spacer(1, 3*mm))

    sub_tbl = Table(
        [[Paragraph(f"Periode: {b_nama} {t_nama}   |   Divisi: {divisi_filter}", sub_style)]],
        colWidths=[267*mm]
    )
    sub_tbl.setStyle(TableStyle([
        ("BACKGROUND",    (0,0), (-1,-1), GREY_L),
        ("TOPPADDING",    (0,0), (-1,-1), 4),
        ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    elems.append(sub_tbl)
    elems.append(Spacer(1, 4*mm))

    # Tabel data
    col_w    = [10*mm, 65*mm, 45*mm, 22*mm, 26*mm, 28*mm, 26*mm]
    header   = ["No","Nama Karyawan","Divisi","Hadir","Terlambat","Tidak Hadir","Total Jam"]
    tbl_data = [header]
    for i, row in enumerate(rows, 1):
        tbl_data.append([
            str(i), row[0], row[1],
            str(row[2]), str(row[3]), str(row[4]),
            f"{row[5]} jam"
        ])
    # Baris total
    tbl_data.append([
        "TOTAL", "", "",
        str(sum(r[2] for r in rows)),
        str(sum(r[3] for r in rows)),
        str(sum(r[4] for r in rows)),
        f"{round(sum(r[5] for r in rows), 1)} jam"
    ])

    tbl = Table(tbl_data, colWidths=col_w, repeatRows=1)
    style_cmds = [
        ("BACKGROUND",     (0,0),  (-1,0),  GREEN),
        ("TEXTCOLOR",      (0,0),  (-1,0),  WHITE),
        ("FONTNAME",       (0,0),  (-1,0),  "Helvetica-Bold"),
        ("FONTSIZE",       (0,0),  (-1,-1), 8.5),
        ("ALIGN",          (0,0),  (-1,-1), "CENTER"),
        ("ALIGN",          (1,1),  (2,-2),  "LEFT"),
        ("VALIGN",         (0,0),  (-1,-1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0,1),  (-1,-2), [WHITE, GREEN_L]),
        ("TOPPADDING",     (0,0),  (-1,-1), 4),
        ("BOTTOMPADDING",  (0,0),  (-1,-1), 4),
        ("GRID",           (0,0),  (-1,-1), 0.4, colors.HexColor("#E5E7EB")),
        ("BACKGROUND",     (0,-1), (-1,-1), GREEN),
        ("TEXTCOLOR",      (0,-1), (-1,-1), WHITE),
        ("FONTNAME",       (0,-1), (-1,-1), "Helvetica-Bold"),
        ("SPAN",           (0,-1), (2,-1)),
    ]
    for ri, row in enumerate(rows, 1):
        if row[3] > 0:
            style_cmds += [
                ("BACKGROUND", (4,ri), (4,ri), YELLOW),
                ("TEXTCOLOR",  (4,ri), (4,ri), colors.HexColor("#A16207")),
            ]
        if row[4] > 0:
            style_cmds += [
                ("BACKGROUND", (5,ri), (5,ri), RED_L),
                ("TEXTCOLOR",  (5,ri), (5,ri), colors.HexColor("#DC3545")),
            ]
    tbl.setStyle(TableStyle(style_cmds))
    elems.append(tbl)

    doc.build(elems)
    buf.seek(0)
    fname = urllib.parse.quote(f"Rekap_Absensi_{b_nama}_{t_nama}.pdf")
    resp  = make_response(buf.read())
    resp.headers["Content-Type"]        = "application/pdf"
    resp.headers["Content-Disposition"] = f"attachment; filename*=UTF-8''{fname}"
    return resp
